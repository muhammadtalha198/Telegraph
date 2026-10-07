#!/usr/bin/env python3
"""Generate one miner YAML per Intent Catalog row into realIntentYamls/."""
from __future__ import annotations

import json
import re
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
MC = ROOT / "MinerCreator"
OUT = MC / "realIntentYamls"
CATALOG_XLSX = ROOT / "Intent Catalog v2 25-Sept.xlsx"
FREE_JSON = MC / "out" / "free-apis-catalog.json"
ID_START = 20001

LLM_INTENTS = {
    "CODE_REVIEW",
    "LLM_OUTPUT_EVALUATION",
    "CONTRACT_OBLIGATION_AUDIT",
    "TASK_EXECUTION_QUALITY",
    "TEXT_SUMMARIZATION",
    "QUESTION_ANSWERING",
    "CHATBOT_CONVERSATION",
    "CODE_GENERATION",
    "CHAT_COMPLETION",
    "MULTIMODAL_INFERENCE",
    "CUSTOMER_TICKET_RESOLUTION",
    "PLAGIARISM_DETECTION",
}

OVERRIDES = {
    "FX_NOW": {
        "method": "GET",
        "base_url": "https://api.frankfurter.app",
        "external_path": "/latest",
        "query": ["from", "to"],
        "path": [],
    },
    "CURRENCY_EXCHANGE": {
        "method": "GET",
        "base_url": "https://api.frankfurter.app",
        "external_path": "/latest",
        "query": ["from", "to", "amount"],
        "path": [],
    },
    "CRYPTO_PRICE_LOOKUP": {
        "method": "GET",
        "base_url": "https://api.coingecko.com/api/v3",
        "external_path": "/simple/price",
        "query": ["ids", "vs_currencies"],
        "path": [],
    },
    "WEATHER_CURRENT": {
        "method": "GET",
        "base_url": "https://api.open-meteo.com",
        "external_path": "/v1/forecast",
        "query": ["latitude", "longitude", "current"],
        "path": [],
    },
    "WEATHER_FORECAST_VERIFY": {
        "method": "GET",
        "base_url": "https://api.open-meteo.com",
        "external_path": "/v1/forecast",
        "query": ["latitude", "longitude", "hourly"],
        "path": [],
    },
    "CRYPTO_YIELD_RATE": {
        "method": "GET",
        "base_url": "https://yields.llama.fi",
        "external_path": "/pools",
        "query": [],
        "path": [],
    },
    "URL_CONTENT_EXTRACTION": {
        "method": "GET",
        "base_url": "https://r.jina.ai",
        "external_path": "/{url}",
        "query": [],
        "path": ["url"],
    },
    "DNS_RECORD_LOOKUP": {
        "method": "GET",
        "base_url": "https://cloudflare-dns.com",
        "external_path": "/dns-query",
        "query": ["name", "type"],
        "path": [],
    },
    "SSL_CERTIFICATE_VERIFY": {
        "method": "GET",
        "base_url": "https://crt.sh",
        "external_path": "/",
        "query": ["q", "output"],
        "path": [],
    },
    "EVENT_OUTCOME_RESOLUTION": {
        "method": "GET",
        "base_url": "https://gamma-api.polymarket.com",
        "external_path": "/events",
        "query": ["closed", "limit"],
        "path": [],
    },
    "LIQUIDITY_DEPTH_VERIFY": {
        "method": "GET",
        "base_url": "https://api.geckoterminal.com",
        "external_path": "/api/v2/networks/{network}/pools/{pool}",
        "query": [],
        "path": ["network", "pool"],
    },
    "WASH_TRADING_DETECTION": {
        "method": "GET",
        "base_url": "https://api.dexscreener.com",
        "external_path": "/latest/dex/tokens/{tokenAddress}",
        "query": [],
        "path": ["tokenAddress"],
    },
    "OPTIMAL_EXECUTION_ROUTE": {
        "method": "GET",
        "base_url": "https://li.quest",
        "external_path": "/v1/quote",
        "query": ["fromChain", "toChain", "fromToken", "toToken", "fromAmount"],
        "path": [],
    },
    "MACRO_ECONOMIC_INDICATOR": {
        "method": "GET",
        "base_url": "https://api.stlouisfed.org",
        "external_path": "/fred/series/observations",
        "query": ["series_id", "api_key", "file_type"],
        "path": [],
    },
    "REGULATORY_FILING_MONITOR": {
        "method": "GET",
        "base_url": "https://data.sec.gov",
        "external_path": "/submissions/CIK{cik}.json",
        "query": [],
        "path": ["cik"],
    },
    "GRAMMAR_SPELL_CHECK": {
        "method": "POST",
        "base_url": "https://api.languagetool.org",
        "external_path": "/v2/check",
        "query": [],
        "path": [],
    },
    "LANGUAGE_TRANSLATION": {
        "method": "POST",
        "base_url": "https://libretranslate.com",
        "external_path": "/translate",
        "query": [],
        "path": [],
    },
    "PORT_SCAN_AUDIT": {
        "method": "GET",
        "base_url": "https://internetdb.shodan.io",
        "external_path": "/{ip}",
        "query": [],
        "path": ["ip"],
    },
    "TOKEN_TOTAL_SUPPLY_VERIFY": {
        "method": "GET",
        "base_url": "https://api.etherscan.io",
        "external_path": "/api",
        "query": ["module", "action", "contractaddress", "apikey"],
        "path": [],
    },
    "GAS_PRICE_ESTIMATION": {
        "method": "GET",
        "base_url": "https://api.etherscan.io",
        "external_path": "/api",
        "query": ["module", "action", "apikey"],
        "path": [],
    },
    "ONCHAIN_METRIC_VERIFY": {
        "method": "GET",
        "base_url": "https://api.etherscan.io",
        "external_path": "/api",
        "query": ["module", "action", "address", "apikey"],
        "path": [],
    },
    "CRYPTO_TRANSFER_VERIFY": {
        "method": "GET",
        "base_url": "https://api.etherscan.io",
        "external_path": "/api",
        "query": ["module", "action", "txhash", "apikey"],
        "path": [],
    },
    "STOCK_PRICE_QUOTE": {
        "method": "GET",
        "base_url": "https://finnhub.io",
        "external_path": "/api/v1/quote",
        "query": ["symbol", "token"],
        "path": [],
    },
    "AIR_QUALITY_INDEX": {
        "method": "GET",
        "base_url": "https://api.openaq.org",
        "external_path": "/v3/locations",
        "query": ["coordinates", "radius"],
        "path": [],
    },
    "THREAT_IP_REPUTATION": {
        "method": "GET",
        "base_url": "https://api.abuseipdb.com",
        "external_path": "/api/v2/check",
        "query": ["ipAddress", "maxAgeInDays"],
        "path": [],
    },
    "SECURITY_REVIEW": {
        "method": "POST",
        "base_url": "https://api.osv.dev",
        "external_path": "/v1/query",
        "query": [],
        "path": [],
    },
    "VULNERABILITY_TRIAGE": {
        "method": "GET",
        "base_url": "https://services.nvd.nist.gov",
        "external_path": "/rest/json/cves/2.0",
        "query": ["cveId"],
        "path": [],
    },
    "VALIDATOR_PERFORMANCE_VERIFY": {
        "method": "GET",
        "base_url": "https://beaconcha.in",
        "external_path": "/api/v1/validator/{index}",
        "query": [],
        "path": ["index"],
    },
    "SENSOR_TELEMETRY_VERIFY": {
        "method": "GET",
        "base_url": "https://api.opensensemap.org",
        "external_path": "/boxes/{senseBoxId}",
        "query": [],
        "path": ["senseBoxId"],
    },
    "ROUTE_ETA": {
        "method": "POST",
        "base_url": "https://api.openrouteservice.org",
        "external_path": "/v2/directions/driving-car",
        "query": [],
        "path": [],
    },
    "FACT_CHECKING": {
        "method": "GET",
        "base_url": "https://factchecktools.googleapis.com",
        "external_path": "/v1alpha1/claims:search",
        "query": ["query", "key"],
        "path": [],
    },
    "TOXICITY_MODERATION": {
        "method": "POST",
        "base_url": "https://commentanalyzer.googleapis.com",
        "external_path": "/v1alpha1/comments:analyze",
        "query": ["key"],
        "path": [],
    },
    "URL_SAFE": {
        "method": "POST",
        "base_url": "https://safebrowsing.googleapis.com",
        "external_path": "/v4/threatMatches:find",
        "query": ["key"],
        "path": [],
    },
    "MALWARE_DETECTION": {
        "method": "GET",
        "base_url": "https://www.virustotal.com",
        "external_path": "/api/v3/files/{hash}",
        "query": [],
        "path": ["hash"],
    },
    "THREAT_INTELLIGENCE": {
        "method": "GET",
        "base_url": "https://otx.alienvault.com",
        "external_path": "/api/v1/indicators/IPv4/{ip}/general",
        "query": [],
        "path": ["ip"],
    },
    "SANCTIONS_SCREENING_MATCH": {
        "method": "GET",
        "base_url": "https://api.trade.gov",
        "external_path": "/consolidated_screening_list/search",
        "query": ["name", "subscription-key"],
        "path": [],
    },
    "VENDOR_VERIFY": {
        "method": "GET",
        "base_url": "https://api.trade.gov",
        "external_path": "/consolidated_screening_list/search",
        "query": ["name", "subscription-key"],
        "path": [],
    },
    "CORPORATE_REGISTRY_LOOKUP": {
        "method": "GET",
        "base_url": "https://api.company-information.service.gov.uk",
        "external_path": "/company/{number}",
        "query": [],
        "path": ["number"],
    },
    "WEB_SEARCH_QUERY": {
        "method": "GET",
        "base_url": "https://api.duckduckgo.com",
        "external_path": "/",
        "query": ["q", "format"],
        "path": [],
    },
}


def slugify(intent: str) -> str:
    s = intent.lower().replace("_", "-")
    s = re.sub(r"[^a-z0-9-]", "", s)
    return re.sub(r"-+", "-", s).strip("-")[:60] or "intent"


def yaml_escape(s: str) -> str:
    return (s or "").replace('"', "'").replace("\n", " ").strip()


def auth_block(auth: str) -> str:
    if auth in (None, "", "none", "self_host"):
        return "auth:\n  type: none"
    return "auth:\n  type: bearer"


def parse_hint(hint: str):
    if not hint:
        return None
    method = "GET"
    m = re.search(r"\b(GET|POST|PUT|DELETE)\s+(https?://\S+)", hint, re.I)
    if m:
        method = m.group(1).upper()
        url = m.group(2)
    else:
        m2 = re.search(r"(https?://\S+)", hint)
        if not m2:
            return None
        url = m2.group(1)
        if "POST" in hint.upper():
            method = "POST"
    url = url.rstrip(").,;")
    url = url.split()[0].split("|")[0]
    placeholders = re.findall(r"\{([^}]+)\}", url)
    for i, p in enumerate(placeholders):
        url = url.replace("{" + p + "}", f"__P{i}__")
    u = urlparse(url)
    path = u.path or "/"
    for i, p in enumerate(placeholders):
        path = path.replace(f"__P{i}__", "{" + p + "}")
    q = parse_qs(u.query, keep_blank_values=True)
    query = list(q.keys())
    path_params = re.findall(r"\{([^}]+)\}", path)
    return {
        "method": method,
        "base_url": f"{u.scheme}://{u.netloc}",
        "external_path": path if path.startswith("/") else "/" + path,
        "query": query,
        "path": path_params,
    }


def params_yaml(path_names, query_names, method: str) -> str:
    parts = []
    if path_names:
        lines = ["      path:", "        required:"]
        for n in path_names:
            lines += [
                f"          - name: {n}",
                "            type: string",
                '            intents: ["*"]',
                f"            description: Path parameter {n}",
            ]
        parts.append("\n".join(lines))
    if query_names:
        lines = ["      query:", "        required:"]
        for n in query_names:
            lines += [
                f"          - name: {n}",
                "            type: string",
                '            intents: ["*"]',
                f"            description: Query parameter {n}",
            ]
        parts.append("\n".join(lines))
    if method == "POST" and not path_names and not query_names:
        parts.append(
            "\n".join(
                [
                    "      body:",
                    "        required:",
                    "          - name: input",
                    "            type: string",
                    '            intents: ["*"]',
                    "            description: Primary request payload / text input",
                ]
            )
        )
    if not parts:
        parts.append(
            "\n".join(
                [
                    "      query:",
                    "        optional:",
                    "          - name: q",
                    "            type: string",
                    '            intents: ["*"]',
                    "            description: Optional query string",
                ]
            )
        )
    return "\n".join(parts)


def render_llm(idx, intent, cat, desc, api) -> str:
    slug = slugify(intent)
    name = f"{intent.replace('_', ' ').title()} Miner"
    docs = api.get("docs") or "https://console.groq.com/docs"
    notes = api.get("notes") or ""
    model = "llama-3.1-8b-instant"
    return f'''version: "1"
kind: miner
id: {idx}
slug: {slug}
protocol: generic
name: {name}
description: >
  {yaml_escape(desc) or yaml_escape(notes) or name}.
  Free-tier LLM backend ({yaml_escape(api.get("api", ""))}). Catalog category: {yaml_escape(cat)}.

base_url: https://api.groq.com/openai

{auth_block(api.get("auth", "free_key"))}

rate_limit_per_sec: 2
cache_ttl_sec: 0
circuit_threshold: 5
circuit_cooldown_seconds: 30

docs:
  documentation: {docs}

endpoints:
  - path: /run
    external_path: /v1/chat/completions
    method: POST
    description: OpenAI-compatible chat completion for intent {intent}.
    intents: [{intent}]
    params:
      body:
        required:
          - name: model
            type: string
            intents: ["*"]
            description: Model id
            accepted_fields:
              "{model}":
                description: Free-tier Groq model
                intents: [{intent}]
          - name: messages
            type: array
            item_type: object
            intents: ["*"]
            description: Chat turns [{{role, content}}]
        optional:
          - name: max_tokens
            type: integer
            intents: ["*"]
          - name: temperature
            type: number
            intents: ["*"]

semantics:
  signal_mapping:
    label_field: choices
  supported_intents:
    - {intent}

on_chain:
  description: Stores model reply for {intent}.
  transform: direct
  min_price_usdc: 0.01
  fields:
    strings:
      - index: 0
        name: response_text
        description: Assistant reply text
        source_path: choices.0.message.content
      - index: 1
        name: model
        description: Model used
        source_path: model
    integers:
      - index: 0
        name: completion_tokens
        description: Completion token count
        source_path: usage.completion_tokens
      - index: 1
        name: prompt_tokens
        description: Prompt token count
        source_path: usage.prompt_tokens
'''


def render_http(idx, intent, cat, desc, api, parsed) -> str:
    slug = slugify(intent)
    name = f"{intent.replace('_', ' ').title()} Miner"
    docs = api.get("docs") or "n/a"
    notes = api.get("notes") or ""
    method = parsed["method"]
    params = params_yaml(parsed.get("path") or [], parsed.get("query") or [], method)
    return f'''version: "1"
kind: miner
id: {idx}
slug: {slug}
protocol: generic
name: {name}
description: >
  {yaml_escape(desc) or yaml_escape(notes) or name}.
  Free API: {yaml_escape(api.get("api", ""))} ({api.get("status", "")}/{api.get("confidence", "")}).
  Catalog category: {yaml_escape(cat)}. Notes: {yaml_escape(notes)}

base_url: {parsed["base_url"]}

{auth_block(api.get("auth", "none"))}

rate_limit_per_sec: 2
cache_ttl_sec: 30
circuit_threshold: 5
circuit_cooldown_seconds: 30

docs:
  documentation: {docs}

endpoints:
  - path: /query
    external_path: {parsed["external_path"]}
    method: {method}
    description: Upstream call for intent {intent} via {yaml_escape(api.get("api", ""))}.
    intents: [{intent}]
    params:
{params}

semantics:
  signal_mapping:
    label_field: data
  supported_intents:
    - {intent}

on_chain:
  description: Stores primary string signal for {intent}.
  transform: direct
  min_price_usdc: 0.01
  fields:
    strings:
      - index: 0
        name: result
        description: Primary result payload or summary field
        source_path: result
      - index: 1
        name: intent
        description: Intent name
        source_path: intent
    integers:
      - index: 0
        name: ok
        description: 1 if upstream succeeded
        source_path: ok
'''


def render_stub(idx, intent, cat, desc, api) -> str:
    slug = slugify(intent)
    name = f"{intent.replace('_', ' ').title()} Miner"
    notes = api.get("notes") or "No free public API found — placeholder for later wiring."
    docs = api.get("docs") or "n/a"
    return f'''version: "1"
kind: miner
id: {idx}
slug: {slug}
protocol: generic
name: {name}
description: >
  {yaml_escape(desc) or name}.
  STATUS: {api.get("status", "not_found")}. {yaml_escape(notes)}
  Catalog category: {yaml_escape(cat)}. Placeholder YAML — replace base_url/endpoint before register.

base_url: https://example.invalid

auth:
  type: none

rate_limit_per_sec: 1
cache_ttl_sec: 0
circuit_threshold: 5
circuit_cooldown_seconds: 30

docs:
  documentation: {docs}

endpoints:
  - path: /query
    external_path: /
    method: GET
    description: Placeholder endpoint for {intent} — not registered yet.
    intents: [{intent}]
    params:
      query:
        optional:
          - name: q
            type: string
            intents: ["*"]
            description: Placeholder query

semantics:
  signal_mapping:
    label_field: data
  supported_intents:
    - {intent}

on_chain:
  description: Placeholder on-chain fields for {intent}.
  transform: direct
  min_price_usdc: 0.01
  fields:
    strings:
      - index: 0
        name: result
        description: Placeholder
        source_path: result
'''


def load_catalog():
    wb = openpyxl.load_workbook(CATALOG_XLSX, data_only=True)
    ws = wb.active
    rows = []
    for r in range(2, ws.max_row + 1):
        intent = ws.cell(r, 3).value
        if not (intent and isinstance(intent, str) and "_" in intent):
            continue
        rows.append(
            {
                "category": ws.cell(r, 2).value or "",
                "intent": intent.strip(),
                "description": (ws.cell(r, 9).value or "").strip(),
            }
        )
    return rows


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    for f in OUT.glob("*.yaml"):
        f.unlink()
    for f in ("README.md", "manifest.json"):
        p = OUT / f
        if p.exists():
            p.unlink()

    catalog = load_catalog()
    apis = {x["intent"]: x for x in json.loads(FREE_JSON.read_text())}
    written = []

    for i, row in enumerate(catalog):
        intent = row["intent"]
        idx = ID_START + i
        api = apis.get(
            intent,
            {
                "api": "unknown",
                "docs": "",
                "auth": "none",
                "endpoint_hint": "",
                "notes": "missing from free-api map",
                "confidence": "low",
                "status": "not_found",
            },
        )

        use_llm = intent in LLM_INTENTS
        if intent == "SECURITY_REVIEW":
            use_llm = False

        if use_llm:
            text = render_llm(idx, intent, row["category"], row["description"], api)
            kind = "llm"
        elif intent in OVERRIDES:
            text = render_http(idx, intent, row["category"], row["description"], api, OVERRIDES[intent])
            kind = "http"
        else:
            parsed = parse_hint(api.get("endpoint_hint") or "")
            if (
                parsed
                and parsed["base_url"].startswith("http")
                and "localhost" not in parsed["base_url"]
                and api.get("status") != "not_found"
            ):
                text = render_http(idx, intent, row["category"], row["description"], api, parsed)
                kind = "http"
            else:
                text = render_stub(idx, intent, row["category"], row["description"], api)
                kind = "stub"

        fname = f"{slugify(intent)}.yaml"
        (OUT / fname).write_text(text)
        written.append(
            {
                "intent": intent,
                "file": fname,
                "id": idx,
                "api": api.get("api"),
                "status": api.get("status"),
                "kind": kind,
            }
        )

    lines = [
        "# realIntentYamls",
        "",
        f"Total: **{len(written)}** YAML files (one per Intent Catalog intent).",
        "",
        "Register later. Bearer keys / self-host URLs may still need wiring.",
        "",
        "| # | Intent | File | YAML id | Kind | Free API | Status |",
        "|---|--------|------|---------|------|----------|--------|",
    ]
    for n, w in enumerate(written, 1):
        lines.append(
            f"| {n} | `{w['intent']}` | `{w['file']}` | {w['id']} | {w['kind']} | {w['api']} | {w['status']} |"
        )
    (OUT / "README.md").write_text("\n".join(lines) + "\n")
    (OUT / "manifest.json").write_text(json.dumps(written, indent=2))

    stubs = sum(1 for w in written if w["kind"] == "stub")
    llms = sum(1 for w in written if w["kind"] == "llm")
    https = sum(1 for w in written if w["kind"] == "http")
    print(f"wrote {len(written)} yamls -> {OUT}")
    print(f"http={https} llm={llms} stub={stubs}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
